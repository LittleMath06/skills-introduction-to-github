"""Importação, validação e normalização da base de clientes atuais (CSV ou XLSX)."""
from __future__ import annotations

import csv
import io
import re
import uuid
from dataclasses import dataclass, field

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..domain import cnpj as cnpj_mod
from ..domain.regions import normalize_uf
from ..domain.text import clean_name, norm, only_digits
from ..models import Company, Customer

MAX_ROWS = 50_000
MAX_BYTES = 20 * 1024 * 1024

# Sinônimos de cabeçalho (normalizados: minúsculo, sem acento, sem pontuação)
HEADER_SYNONYMS = {
    "cnpj": ["cnpj", "cnpj cpf", "cpf cnpj", "documento", "cnpj do cliente", "cnpj cliente", "doc"],
    "razao_social": ["razao social", "razao", "nome empresarial", "cliente", "empresa", "nome",
                     "nome do cliente"],
    "nome_fantasia": ["nome fantasia", "fantasia"],
    "municipio": ["cidade", "municipio", "localidade"],
    "uf": ["uf", "estado", "sigla uf"],
    "cnae": ["cnae", "cnae principal", "codigo cnae", "cnae fiscal"],
    "segmento": ["segmento", "ramo", "setor", "categoria", "tipo de cliente", "atividade"],
    "porte": ["porte"],
}
_PORTE_TEXT = {"me": "01", "micro": "01", "micro empresa": "01", "microempresa": "01",
               "epp": "03", "pequeno porte": "03", "empresa de pequeno porte": "03",
               "demais": "05", "grande": "05", "media": "05"}


@dataclass
class ImportReport:
    batch: str
    total_rows: int = 0
    imported: int = 0
    updated: int = 0
    duplicates_in_file: int = 0
    invalid_cnpj: list[dict] = field(default_factory=list)
    without_cnpj: int = 0
    linked: int = 0
    not_found: int = 0
    ignored_empty: int = 0
    columns: dict[str, str] = field(default_factory=dict)
    unmapped_columns: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {**self.__dict__, "invalid_cnpj": self.invalid_cnpj[:200],
                "invalid_cnpj_count": len(self.invalid_cnpj)}


class ImportError_(ValueError):
    pass


def _norm_header(h: str) -> str:
    return re.sub(r"[^a-z0-9 ]", " ", norm(h)).strip()


def map_headers(headers: list[str]) -> tuple[dict[str, int], list[str]]:
    normalized = [re.sub(r"\s+", " ", _norm_header(h or "")) for h in headers]
    mapping: dict[str, int] = {}
    for fieldname, synonyms in HEADER_SYNONYMS.items():
        for syn in synonyms:  # ordem de preferência
            if syn in normalized and normalized.index(syn) not in mapping.values():
                mapping[fieldname] = normalized.index(syn)
                break
    unmapped = [headers[i] for i in range(len(headers)) if i not in mapping.values() and headers[i]]
    return mapping, unmapped


def read_table(filename: str, content: bytes) -> list[list[str]]:
    if len(content) > MAX_BYTES:
        raise ImportError_("Arquivo maior que 20 MB")
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # arquivo corrompido ou não-XLSX
            raise ImportError_("Não foi possível ler a planilha XLSX") from exc
        ws = wb.worksheets[0]
        rows = [["" if v is None else str(v) for v in row]
                for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    if name.endswith((".csv", ".txt")):
        for enc in ("utf-8-sig", "latin-1"):
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        sample = text[:5000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    raise ImportError_("Formato não suportado. Envie .csv ou .xlsx")


def _find_header_row(rows: list[list[str]]) -> int:
    for i, row in enumerate(rows[:20]):
        mapping, _ = map_headers(row)
        if "cnpj" in mapping or ("razao_social" in mapping and len(mapping) >= 2):
            return i
    raise ImportError_("Cabeçalho não reconhecido: é necessária ao menos a coluna CNPJ "
                       "(ou razão social + cidade/UF/CNAE)")


def import_customers(session: Session, filename: str, content: bytes) -> ImportReport:
    rows = read_table(filename, content)
    header_idx = _find_header_row(rows)
    headers = rows[header_idx]
    mapping, unmapped = map_headers(headers)
    report = ImportReport(batch=uuid.uuid4().hex[:12])
    report.columns = {k: headers[v] for k, v in mapping.items()}
    report.unmapped_columns = unmapped
    data_rows = rows[header_idx + 1:]
    if len(data_rows) > MAX_ROWS:
        raise ImportError_(f"Máximo de {MAX_ROWS} linhas por importação")

    def cell(row: list[str], fieldname: str) -> str | None:
        idx = mapping.get(fieldname)
        if idx is None or idx >= len(row):
            return None
        v = (row[idx] or "").strip()
        return v or None

    seen: set[str] = set()
    for line_no, row in enumerate(data_rows, start=header_idx + 2):
        if not any((c or "").strip() for c in row):
            report.ignored_empty += 1
            continue
        report.total_rows += 1
        raw_cnpj = cell(row, "cnpj")
        cnpj = None
        if raw_cnpj:
            cleaned = cnpj_mod.clean(raw_cnpj)
            if not cnpj_mod.is_valid(cleaned):
                # 11 dígitos que não formam CNPJ válido: provavelmente CPF (pessoa física)
                is_cpf = len(only_digits(raw_cnpj)) == 11
                report.invalid_cnpj.append({
                    "linha": line_no,
                    "valor": f"{raw_cnpj} (CPF/pessoa física — ignorado)" if is_cpf else raw_cnpj,
                })
                continue
            cnpj = cleaned
            if cnpj in seen:
                report.duplicates_in_file += 1
                continue
            seen.add(cnpj)
        cnae = only_digits(cell(row, "cnae") or "")
        cnae = cnae.zfill(7) if 5 <= len(cnae) <= 7 else None
        porte_raw = norm(cell(row, "porte"))
        values = {
            "razao_social": clean_name(cell(row, "razao_social")),
            "nome_fantasia": clean_name(cell(row, "nome_fantasia")),
            "municipio": clean_name(cell(row, "municipio")),
            "uf": normalize_uf(cell(row, "uf")),
            "cnae": cnae,
            "segmento_informado": clean_name(cell(row, "segmento")),
            "porte": porte_raw if porte_raw in {"00", "01", "03", "05"} else _PORTE_TEXT.get(porte_raw),
            "raw": {headers[i]: (row[i] if i < len(row) else None)
                    for i in range(len(headers)) if headers[i]},
        }
        if not cnpj and not (values["cnae"] or values["uf"] or values["razao_social"]):
            report.ignored_empty += 1
            continue
        customer = session.scalar(select(Customer).where(Customer.cnpj == cnpj)) if cnpj else None
        if customer is None:
            customer = Customer(cnpj=cnpj, import_batch=report.batch, **values)
            session.add(customer)
            report.imported += 1
        else:
            for k, v in values.items():
                if v is not None:
                    setattr(customer, k, v)
            customer.import_batch = report.batch
            report.updated += 1
        if not cnpj:
            customer.status = "sem_cnpj"
            report.without_cnpj += 1
            continue
        company = session.scalar(select(Company).where(Company.cnpj == cnpj))
        customer.company_id = company.id if company else None
        customer.status = "ok" if company else "nao_encontrado"
        if company:
            report.linked += 1
        else:
            report.not_found += 1
    session.flush()
    return report


def relink_customers(session: Session) -> int:
    """Liga clientes 'nao_encontrado' a empresas que passaram a existir na base."""
    linked = 0
    for customer in session.scalars(
        select(Customer).where(Customer.status == "nao_encontrado")
    ).all():
        company = session.scalar(select(Company).where(Company.cnpj == customer.cnpj))
        if company:
            customer.company_id = company.id
            customer.status = "ok"
            linked += 1
    return linked
